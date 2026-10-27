#!/usr/bin/env python3
"""Reference image -> editable PowerPoint, full-canvas high-fidelity rebuild.

This is the engine behind the `reference-image-to-editable-pptx` skill.

The governing idea: the reference image's pixel canvas IS the coordinate
system. Every object is measured in reference pixels and mapped straight onto
the slide. There is no contain/fit, no letterbox, no global scale or offset
applied after the mapping. A single `X()`/`Y()` pair converts reference pixels
to slide EMU, and that is the only conversion in the pipeline.

Anything that violates that is the failure mode this engine exists to prevent:
a page that is technically "based on" the reference but sits inside a border
with dead space on the right and bottom.

Pipeline:
    measure -> scene_spec -> build -> render -> compare -> correct -> rebuild

The compare step is not optional. A build that has not been rendered and
diffed against the reference has not been verified at all.

Usage:
    python3 build_editable_ppt.py --reference ref.png --spec scene_spec.json
    python3 build_editable_ppt.py --reference ref.png --spec scene_spec.json \
        --out ppt/ --render preview/
"""
from __future__ import annotations

import argparse
import json
import math
import os
import shutil
import subprocess
import sys
import zipfile
from pathlib import Path

try:
    import numpy as np
    from PIL import Image, ImageChops
except ImportError:
    sys.exit("ERROR: requires Pillow and NumPy")

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.oxml.ns import qn
    from pptx.oxml.xmlchemy import OxmlElement
    from pptx.util import Emu, Inches, Pt
except ImportError:
    sys.exit("ERROR: requires python-pptx")


# ---------------------------------------------------------------------------
# Canvas and coordinate mapping
# ---------------------------------------------------------------------------

class Canvas:
    """The reference image's pixel canvas, mapped 1:1 onto the slide.

    `X(0)` is the slide's left edge and `X(REF_W)` is its right edge. There is
    no uniform_scale, no offset_x, no offset_y. That is deliberate: every
    known failure of this task traces back to one of those three.

    All arithmetic happens in integer EMU rather than float inches. Converting
    through `Inches(float)` truncates, and a truncated slide height makes the
    page aspect ratio differ from the reference by a few parts per million —
    small, but it means the "aspect ratio error: 0" acceptance check is only
    ever true by luck. Computing the slide size once in EMU and dividing the
    reference pixel span against it makes the ratio exact by construction.
    """

    EMU_PER_INCH = 914400

    def __init__(self, ref_w: int, ref_h: int, slide_w_in: float = 16.0):
        self.ref_w = ref_w
        self.ref_h = ref_h

        # Exact aspect ratio: reduce the reference to lowest terms and build
        # the slide as an integer multiple of that reduced pair. A plain
        # round(slide_w * h / w) leaves the ratio off by a fraction of an EMU,
        # which is invisible in a render but makes the "aspect ratio error: 0"
        # acceptance check literally untrue.
        g = math.gcd(ref_w, ref_h)
        self.rw = ref_w // g
        self.rh = ref_h // g

        target_w_emu = int(round(slide_w_in * self.EMU_PER_INCH))
        k = max(1, round(target_w_emu / self.rw))
        self.slide_w_emu = k * self.rw
        self.slide_h_emu = k * self.rh

        self.slide_w_in = self.slide_w_emu / self.EMU_PER_INCH
        self.slide_h_in = self.slide_h_emu / self.EMU_PER_INCH
        self.ratio = ref_w / ref_h

    def _emu(self, value: float, span_px: int, span_emu: int) -> int:
        return int(round(value / span_px * span_emu))

    def X(self, px: float):
        return Emu(self._emu(px, self.ref_w, self.slide_w_emu))

    def Y(self, py: float):
        return Emu(self._emu(py, self.ref_h, self.slide_h_emu))

    def W(self, pw: float):
        return Emu(self._emu(pw, self.ref_w, self.slide_w_emu))

    def H(self, ph: float):
        return Emu(self._emu(ph, self.ref_h, self.slide_h_emu))

    def as_dict(self) -> dict:
        return {
            "reference_size_px": [self.ref_w, self.ref_h],
            "reference_ratio": round(self.ratio, 6),
            "slide_size_in": [round(self.slide_w_in, 4),
                              round(self.slide_h_in, 4)],
            "slide_size_emu": [self.slide_w_emu, self.slide_h_emu],
            "slide_ratio": round(self.slide_w_emu / self.slide_h_emu, 6),
        }


def load_reference(path: Path):
    """Read the true pixel size from the file, never from a thumbnail."""
    im = Image.open(path).convert("RGB")
    return im, im.size[0], im.size[1]


# ---------------------------------------------------------------------------
# Icon handling
# ---------------------------------------------------------------------------

def svg_to_png_128(svg_path: Path, png_path: Path) -> bool:
    """Render an SVG to a 128x128 transparent PNG.

    CairoSVG is the specified tool, but it needs libcairo, which is not
    present on every machine. rsvg-convert (librsvg) does the same job with no
    cairo dependency, and Quick Look is the last resort. All three produce a
    128x128 RGBA PNG, so the rest of the pipeline is identical either way.
    """
    png_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        import cairosvg
        cairosvg.svg2png(url=str(svg_path), write_to=str(png_path),
                         output_width=128, output_height=128)
        return True
    except Exception:
        pass

    rsvg = shutil.which("rsvg-convert")
    if rsvg:
        try:
            subprocess.run([rsvg, "-w", "128", "-h", "128",
                            "-o", str(png_path), str(svg_path)],
                           check=True, capture_output=True, timeout=60)
            return True
        except Exception:
            pass

    ql = shutil.which("qlmanage")
    if ql:
        try:
            tmp = png_path.parent / "_ql"
            tmp.mkdir(exist_ok=True)
            subprocess.run([ql, "-t", "-s", "128", "-o", str(tmp),
                            str(svg_path)], capture_output=True, timeout=60)
            produced = list(tmp.glob("*.png"))
            if produced:
                src = Image.open(produced[0]).convert("RGBA")
                canvas = Image.new("RGBA", (128, 128), (0, 0, 0, 0))
                scale = min(128 / src.width, 128 / src.height)
                nw, nh = int(src.width * scale), int(src.height * scale)
                src = src.resize((nw, nh), Image.Resampling.LANCZOS)
                canvas.paste(src, ((128 - nw) // 2, (128 - nh) // 2), src)
                canvas.save(png_path)
                shutil.rmtree(tmp, ignore_errors=True)
                return True
        except Exception:
            pass
    return False


def alpha_bbox_fraction(png_path: Path):
    """Fractional bounding box of the visible (alpha > 8) pixels."""
    im = Image.open(png_path).convert("RGBA")
    alpha = np.array(im)[:, :, 3]
    ys, xs = np.where(alpha > 8)
    if len(xs) == 0 or len(ys) == 0:
        return 0.0, 0.0, 1.0, 1.0
    w, h = im.size
    x0, x1 = xs.min(), xs.max() + 1
    y0, y1 = ys.min(), ys.max() + 1
    return (x0 / w, y0 / h, (x1 - x0) / w, (y1 - y0) / h)


def place_icon(slide, canvas: Canvas, icon_png: Path, spec: dict):
    """Insert an icon so its VISIBLE outline lands on the target box.

    The PNG is a 128x128 canvas whose content usually does not fill it. Aligning
    the 128x128 frame to the target would visibly misplace the icon, so the
    alpha bbox is measured and the PNG is positioned by its visible content.
    """
    fx, fy, fw, fh = alpha_bbox_fraction(icon_png)
    tx, ty = spec["target_visible_x"], spec["target_visible_y"]
    tw, th = spec["target_visible_w"], spec["target_visible_h"]

    vis_w = fw * 128
    vis_h = fh * 128
    scale = min(tw / vis_w, th / vis_h) if vis_w and vis_h else 1.0
    shape_w = 128 * scale
    shape_h = 128 * scale
    actual_w = vis_w * scale
    actual_h = vis_h * scale
    shape_x = tx + (tw - actual_w) / 2 - fx * 128 * scale
    shape_y = ty + (th - actual_h) / 2 - fy * 128 * scale

    pic = slide.shapes.add_picture(str(icon_png),
                                   canvas.X(shape_x), canvas.Y(shape_y),
                                   canvas.W(shape_w), canvas.H(shape_h))
    pic.name = spec.get("name", f"icon_{spec.get('asset','icon')}")
    return pic


# ---------------------------------------------------------------------------
# Native arrow heads
# ---------------------------------------------------------------------------

def set_native_arrow(connector, end: bool = True, start: bool = False):
    """Attach a native triangle arrowhead to a connector via Open XML.

    The alternatives — Freeform, a PNG arrow, an SVG arrow, a line plus a
    triangle — are all rejected by the acceptance criteria because they are
    not editable connector arrows in PowerPoint's own sense.
    """
    line_xml = connector._element.spPr.get_or_add_ln()
    if end:
        tail = OxmlElement("a:tailEnd")
        tail.set("type", "triangle")
        tail.set("w", "med")
        tail.set("len", "med")
        line_xml.append(tail)
    if start:
        head = OxmlElement("a:headEnd")
        head.set("type", "triangle")
        head.set("w", "med")
        head.set("len", "med")
        line_xml.append(head)


# ---------------------------------------------------------------------------
# Shape builders
# ---------------------------------------------------------------------------

def strip_theme_style(shape):
    """Remove the theme <p:style> reference and any inherited effects.

    python-pptx's add_shape() attaches a <p:style> element that points at the
    theme's shape defaults. In PowerPoint that means a drop shadow, a themed
    fill and a themed line. The reference image has no shadow, so the theme
    style has to go — otherwise every frame renders with a grey halo that the
    reference does not have.
    """
    el = shape._element
    node = el.find(qn("p:style"))
    if node is not None:
        el.remove(node)
    spPr = el.find(qn("p:spPr"))
    if spPr is None:
        return
    effect = spPr.find(qn("a:effectLst"))
    if effect is None:
        effect = OxmlElement("a:effectLst")
        spPr.append(effect)


def px_line_pt(px: float, canvas: "Canvas") -> float:
    """Convert a reference-pixel line width to points.

    A 3px border measured on the reference must stay 3px on the slide. PowerPoint
    stores stroke width in points, so the pixel width is scaled by the same
    factor the canvas uses to map pixels to slide units. Declaring
    `line_width_pt` directly is honoured as-is.
    """
    return px / canvas.ref_w * canvas.slide_w_in * 72.0


def px_font_pt(px: float, canvas: "Canvas") -> float:
    """Convert a reference-pixel font size to points.

    Text measured against the reference is measured in reference pixels, but
    PowerPoint takes points. Converting through the slide width keeps text
    at the same visual scale as the pixels it was measured against. Specs
    that state font_size_pt directly are honoured as-is.
    """
    return px / canvas.ref_w * canvas.slide_w_in * 72.0


def _rgb(hex_str: str):
    h = hex_str.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _apply_line(shape, spec: dict, canvas: "Canvas" = None):
    ln = shape.line
    if spec.get("line"):
        ln.color.rgb = _rgb(spec["line"])
    if spec.get("line_width_px") and canvas is not None:
        ln.width = Pt(px_line_pt(spec["line_width_px"], canvas))
    elif spec.get("line_width_pt"):
        ln.width = Pt(float(spec["line_width_pt"]))
    dash = spec.get("dash")
    if dash:
        mapping = {
            "dash": MSO_LINE_DASH_STYLE.DASH,
            "round_dot": MSO_LINE_DASH_STYLE.ROUND_DOT,
            "square_dot": MSO_LINE_DASH_STYLE.SQUARE_DOT,
            "long_dash": MSO_LINE_DASH_STYLE.LONG_DASH,
            "dash_dot": MSO_LINE_DASH_STYLE.DASH_DOT,
        }
        if dash in mapping:
            ln.dash_style = mapping[dash]


def add_frame(slide, canvas: Canvas, spec: dict):
    """A rounded rectangle, rectangle, ellipse or circle."""
    kind = spec.get("type", "rounded_rectangle")
    shape_type = {
        "rounded_rectangle": MSO_SHAPE.ROUNDED_RECTANGLE,
        "rectangle": MSO_SHAPE.RECTANGLE,
        "oval": MSO_SHAPE.OVAL,
        "circle": MSO_SHAPE.OVAL,
        "diamond": MSO_SHAPE.DIAMOND,
        "chevron": MSO_SHAPE.CHEVRON,
    }.get(kind, MSO_SHAPE.ROUNDED_RECTANGLE)

    shape = slide.shapes.add_shape(
        shape_type, canvas.X(spec["x"]), canvas.Y(spec["y"]),
        canvas.W(spec["w"]), canvas.H(spec["h"]))
    shape.name = spec.get("name", "frame")
    if not spec.get("shadow", False):
        strip_theme_style(shape)

    if spec.get("fill"):
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(spec["fill"])
    else:
        shape.fill.background()

    _apply_line(shape, spec, canvas)

    if kind == "rounded_rectangle" and spec.get("corner_radius"):
        try:
            shape.adjustments[0] = float(spec["corner_radius"])
        except (IndexError, ValueError):
            pass

    if spec.get("text"):
        tf = shape.text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = spec["text"]
        if spec.get("font_size_px"):
            run.font.size = Pt(px_font_pt(spec["font_size_px"], canvas))
        else:
            run.font.size = Pt(spec.get("font_size_pt", 12))
        run.font.bold = spec.get("bold", False)
        run.font.name = spec.get("font", "Microsoft YaHei")
        if spec.get("color"):
            run.font.color.rgb = _rgb(spec["color"])
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = True
    return shape


def add_text(slide, canvas: Canvas, spec: dict):
    """A text box with explicit margins, anchoring and run formatting."""
    box = slide.shapes.add_textbox(
        canvas.X(spec["x"]), canvas.Y(spec["y"]),
        canvas.W(spec["w"]), canvas.H(spec["h"]))
    box.name = spec.get("name", "text")

    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    # Never let PowerPoint shrink or grow the text to fit the box: the size
    # was measured against the reference and must be honoured literally.
    bodyPr = tf._txBody.find(qn("a:bodyPr"))
    for tag in ("a:normAutofit", "a:spAutoFit"):
        node = bodyPr.find(qn(tag))
        if node is not None:
            bodyPr.remove(node)
    bodyPr.append(OxmlElement("a:noAutofit"))
    tf.margin_left = canvas.W(spec.get("margin_left_px", 3))
    tf.margin_right = canvas.W(spec.get("margin_right_px", 3))
    tf.margin_top = canvas.H(spec.get("margin_top_px", 0))
    tf.margin_bottom = canvas.H(spec.get("margin_bottom_px", 0))
    tf.vertical_anchor = {
        "top": MSO_ANCHOR.TOP,
        "middle": MSO_ANCHOR.MIDDLE,
        "bottom": MSO_ANCHOR.BOTTOM,
    }.get(spec.get("vertical_align", "top"), MSO_ANCHOR.TOP)

    lines = spec.get("text", "").split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        p.alignment = {
            "left": PP_ALIGN.LEFT,
            "center": PP_ALIGN.CENTER,
            "right": PP_ALIGN.RIGHT,
        }.get(spec.get("align", "left"), PP_ALIGN.LEFT)
        run = p.add_run()
        run.text = line
        run.font.name = spec.get("font", "Microsoft YaHei")
        if spec.get("font_size_px"):
            run.font.size = Pt(px_font_pt(spec["font_size_px"], canvas))
        else:
            run.font.size = Pt(spec.get("font_size_pt", 12))
        run.font.bold = spec.get("bold", False)
        if spec.get("color"):
            run.font.color.rgb = _rgb(spec["color"])
    return box


def add_connector(slide, canvas: Canvas, spec: dict):
    """A native PowerPoint connector — straight, elbow or curve."""
    ctype = {
        "straight": MSO_CONNECTOR.STRAIGHT,
        "elbow": MSO_CONNECTOR.ELBOW,
        "curve": MSO_CONNECTOR.CURVE,
    }.get(spec.get("type", "straight"), MSO_CONNECTOR.STRAIGHT)

    conn = slide.shapes.add_connector(
        ctype, canvas.X(spec["x1"]), canvas.Y(spec["y1"]),
        canvas.X(spec["x2"]), canvas.Y(spec["y2"]))
    conn.name = spec.get("name", "connector")
    if not spec.get("shadow", False):
        strip_theme_style(conn)
    _apply_line(conn, spec, canvas)
    set_native_arrow(conn, end=spec.get("end_arrow", False),
                     start=spec.get("start_arrow", False))
    return conn


def add_arrow_shape(slide, canvas: Canvas, spec: dict):
    """A native arrow SHAPE (e.g. a bidirectional block arrow)."""
    shape_type = {
        "left_right_arrow": MSO_SHAPE.LEFT_RIGHT_ARROW,
        "right_arrow": MSO_SHAPE.RIGHT_ARROW,
        "down_arrow": MSO_SHAPE.DOWN_ARROW,
        "up_arrow": MSO_SHAPE.UP_ARROW,
    }.get(spec.get("type", "right_arrow"), MSO_SHAPE.RIGHT_ARROW)
    shape = slide.shapes.add_shape(
        shape_type, canvas.X(spec["x"]), canvas.Y(spec["y"]),
        canvas.W(spec["w"]), canvas.H(spec["h"]))
    shape.name = spec.get("name", "arrow_shape")
    if not spec.get("shadow", False):
        strip_theme_style(shape)
    if spec.get("fill"):
        shape.fill.solid()
        shape.fill.fore_color.rgb = _rgb(spec["fill"])
    else:
        shape.fill.background()
    _apply_line(shape, spec, canvas)
    return shape


# ---------------------------------------------------------------------------
# Layer order
# ---------------------------------------------------------------------------

# Later additions sit on top in PowerPoint, so the draw order is the z-order.
DRAW_ORDER = ["background", "connectors_outer", "frames", "headers",
              "cards", "connectors_inner", "icons", "texts", "labels"]


def build_presentation(spec: dict, canvas: Canvas, icon_dir: Path,
                       out_path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Emu(canvas.slide_w_emu)
    prs.slide_height = Emu(canvas.slide_h_emu)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    bg = (spec.get("page") or {}).get("background")
    if bg:
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = _rgb(bg)

    def sorted_by_z(items):
        return sorted(items, key=lambda s: s.get("z", 0))

    frames = spec.get("frames", [])
    texts = spec.get("texts", [])
    icons = spec.get("icons", [])
    connectors = spec.get("connectors", [])

    # Split connectors: "outer" (cross-module) go early so they sit behind
    # the module frames; "inner" go after the cards.
    outer = [c for c in connectors if c.get("layer", "outer") == "outer"]
    inner = [c for c in connectors if c.get("layer", "outer") == "inner"]

    for c in sorted_by_z(outer):
        if c.get("type") in ("left_right_arrow", "right_arrow", "down_arrow",
                             "up_arrow"):
            add_arrow_shape(slide, canvas, c)
        else:
            add_connector(slide, canvas, c)

    for f in sorted_by_z(frames):
        add_frame(slide, canvas, f)

    for c in sorted_by_z(inner):
        if c.get("type") in ("left_right_arrow", "right_arrow", "down_arrow",
                             "up_arrow"):
            add_arrow_shape(slide, canvas, c)
        else:
            add_connector(slide, canvas, c)

    for ic in sorted_by_z(icons):
        png = icon_dir / "png_128" / f"{ic['asset']}.png"
        if not png.exists():
            svg = icon_dir / "svg" / f"{ic['asset']}.svg"
            if svg.exists():
                svg_to_png_128(svg, png)
        if png.exists():
            place_icon(slide, canvas, png, ic)

    for t in sorted_by_z(texts):
        add_text(slide, canvas, t)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


# ---------------------------------------------------------------------------
# Rendering
# ---------------------------------------------------------------------------

def find_soffice():
    for c in ("soffice", "libreoffice"):
        p = shutil.which(c)
        if p:
            return p
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if mac.exists():
        return str(mac)
    return None


def render_pptx(pptx_path: Path, out_dir: Path, prefix: str = "ppt_render",
                dpi: int = 150) -> Path | None:
    """Render the PPTX to a PNG via LibreOffice, then pdftoppm.

    A dedicated LibreOffice user profile is mandatory: reusing the default one
    makes soffice exit 77 when another instance holds the lock.
    """
    soffice = find_soffice()
    if not soffice:
        return None
    out_dir.mkdir(parents=True, exist_ok=True)
    profile = f"/tmp/lo_profile_{prefix}"
    shutil.rmtree(profile, ignore_errors=True)

    cmd = [soffice, f"-env:UserInstallation=file://{profile}",
           "--headless", "--convert-to", "pdf", "--outdir", str(out_dir),
           str(pptx_path)]
    try:
        subprocess.run(cmd, capture_output=True, timeout=300)
    except Exception:
        return None

    pdf = out_dir / (pptx_path.stem + ".pdf")
    if not pdf.exists():
        return None

    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        return None
    subprocess.run([pdftoppm, "-png", "-singlefile", "-r", str(dpi),
                    str(pdf), str(out_dir / prefix)], capture_output=True,
                   timeout=180)
    png = out_dir / f"{prefix}.png"
    return png if png.exists() else None


# ---------------------------------------------------------------------------
# Measurement and comparison
# ---------------------------------------------------------------------------

def content_bbox_from_image(image: Image.Image):
    """Bounding box of meaningful content, ignoring near-white paper."""
    arr = np.array(image.convert("RGB"))
    r, g, b = arr[:, :, 0], arr[:, :, 1], arr[:, :, 2]
    saturation = arr.max(axis=2) - arr.min(axis=2)
    mask = (((saturation > 35) & (arr.min(axis=2) < 245)) |
            ((r < 245) & (g < 245) & (b < 245)))
    ys, xs = np.where(mask)
    if len(xs) == 0:
        return 0, 0, image.width, image.height
    return (int(xs.min()), int(ys.min()),
            int(xs.max() + 1), int(ys.max() + 1))


def bbox_ratios(bbox, width, height):
    x0, y0, x1, y1 = bbox
    return (x0 / width, y0 / height, x1 / width, y1 / height)


def occupancy(bbox, width, height):
    x0, y0, x1, y1 = bbox
    return ((x1 - x0) / width, (y1 - y0) / height)


def bbox_center(bbox, width, height):
    x0, y0, x1, y1 = bbox
    return (((x0 + x1) / 2) / width, ((y0 + y1) / 2) / height)


def make_comparisons(reference: Image.Image, render: Image.Image,
                     out_dir: Path, gap: int = 24,
                     title_height: int = 40) -> dict:
    """Side-by-side, 50% overlay and amplified difference map."""
    out_dir.mkdir(parents=True, exist_ok=True)
    ref = reference.convert("RGB")
    ren = render.convert("RGB").resize(ref.size, Image.Resampling.LANCZOS)

    comparison = Image.new("RGB", (ref.width * 2 + gap, ref.height + title_height),
                           "white")
    comparison.paste(ref, (0, title_height))
    comparison.paste(ren, (ref.width + gap, title_height))
    comparison.save(out_dir / "comparison.png")

    overlay = Image.blend(ref, ren, 0.5)
    overlay.save(out_dir / "overlay_comparison.png")

    diff = ImageChops.difference(ref, ren)
    arr = np.array(diff).astype(np.float32)
    arr = np.clip(arr * 2.5, 0, 255).astype(np.uint8)
    Image.fromarray(arr).save(out_dir / "difference_map.png")

    return {
        "comparison": str(out_dir / "comparison.png"),
        "overlay": str(out_dir / "overlay_comparison.png"),
        "difference": str(out_dir / "difference_map.png"),
    }


def measure_page(reference: Image.Image, render: Image.Image) -> dict:
    ref = reference.convert("RGB")
    ren = render.convert("RGB").resize(ref.size, Image.Resampling.LANCZOS)
    w, h = ref.size
    rb = content_bbox_from_image(ref)
    kb = content_bbox_from_image(ren)
    ro, ko = occupancy(rb, w, h), occupancy(kb, w, h)
    rc, kc = bbox_center(rb, w, h), bbox_center(kb, w, h)
    return {
        "reference_bbox_ratios": [round(v, 6) for v in bbox_ratios(rb, w, h)],
        "render_bbox_ratios": [round(v, 6) for v in bbox_ratios(kb, w, h)],
        "reference_occupancy": [round(v, 6) for v in ro],
        "render_occupancy": [round(v, 6) for v in ko],
        "reference_center": [round(v, 6) for v in rc],
        "render_center": [round(v, 6) for v in kc],
        "occupancy_error_x": round(abs(ko[0] - ro[0]), 6),
        "occupancy_error_y": round(abs(ko[1] - ro[1]), 6),
        "center_error_x": round(abs(kc[0] - rc[0]), 6),
        "center_error_y": round(abs(kc[1] - rc[1]), 6),
    }


# ---------------------------------------------------------------------------
# Out-of-bounds and media checks
# ---------------------------------------------------------------------------

def find_out_of_bounds(slide, presentation):
    warnings = []
    for shape in slide.shapes:
        if (shape.left < 0 or shape.top < 0 or
                shape.left + shape.width > presentation.slide_width or
                shape.top + shape.height > presentation.slide_height):
            warnings.append(shape.name)
    return warnings


def check_media(pptx_path: Path):
    with zipfile.ZipFile(pptx_path) as zf:
        media = [n for n in zf.namelist() if n.startswith("ppt/media/")]
        sizes = {n: zf.getinfo(n).file_size for n in media}
    huge = {n: s for n, s in sizes.items() if s > 4 * 1024 * 1024}
    return {
        "media_count": len(media),
        "media_total_bytes": sum(sizes.values()),
        "huge_media": huge,
        "pptx_bytes": pptx_path.stat().st_size,
    }


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv=None):
    ap = argparse.ArgumentParser(
        description="Rebuild a reference image as an editable PPTX.")
    ap.add_argument("--reference", required=True)
    ap.add_argument("--spec", required=True, help="scene_spec JSON")
    ap.add_argument("--out", default="ppt")
    ap.add_argument("--render", default="preview")
    ap.add_argument("--icons", default="icons")
    ap.add_argument("--out-name", default="editable_reproduction.pptx")
    ap.add_argument("--no-render", action="store_true")
    args = ap.parse_args(argv)

    ref_path = Path(args.reference)
    spec_path = Path(args.spec)
    spec = json.loads(spec_path.read_text(encoding="utf-8"))

    reference, ref_w, ref_h = load_reference(ref_path)
    canvas = Canvas(ref_w, ref_h)
    print(f"reference: {ref_w}x{ref_h}  ratio {canvas.ratio:.4f}")
    print(f"slide: {canvas.slide_w_in:.3f} x {canvas.slide_h_in:.3f} in")

    out_dir = Path(args.out)
    render_dir = Path(args.render)
    icon_dir = Path(args.icons)

    pptx = build_presentation(spec, canvas, icon_dir, out_dir / args.out_name)
    print(f"built: {pptx}")

    report = {"canvas": canvas.as_dict(), "pptx": str(pptx)}

    if not args.no_render:
        png = render_pptx(pptx, render_dir)
        if png:
            render = Image.open(png).convert("RGB")
            report["render"] = str(png)
            report["comparison_files"] = make_comparisons(reference, render,
                                                          render_dir)
            report["measurement"] = measure_page(reference, render)
            m = report["measurement"]
            print(f"occupancy err: x={m['occupancy_error_x']} "
                  f"y={m['occupancy_error_y']}")
            print(f"center err:    x={m['center_error_x']} "
                  f"y={m['center_error_y']}")
        else:
            report["render_error"] = (
                "LibreOffice not found or conversion failed. Install with "
                "`brew install --cask libreoffice`.")
            print("WARNING: no render produced; comparisons skipped.",
                  file=sys.stderr)

    report["media"] = check_media(pptx)
    report_path = render_dir / "report.json"
    report_path.parent.mkdir(parents=True, exist_ok=True)
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2),
                           encoding="utf-8")
    print(f"report: {report_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
