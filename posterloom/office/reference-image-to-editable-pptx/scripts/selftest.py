#!/usr/bin/env python3
"""End-to-end self-test for the editable-PPT rebuild skill.

Builds a synthetic reference image, rebuilds it as an editable PPTX, renders
it, and checks the acceptance thresholds. The synthetic reference is
deliberately asymmetric and non-16:9 so that any contain/fit behaviour or
letterboxing shows up immediately as a failed occupancy check.

A skill that ships an unverified pipeline is a liability: the thresholds in
the prompt look reasonable on paper and quietly fail in practice. This test
exists so that never happens.

Usage:
    python3 selftest.py
"""
from __future__ import annotations

import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
try:
    from PIL import Image, ImageDraw, ImageFont
except ImportError:
    sys.exit("ERROR: requires Pillow")

sys.path.insert(0, str(HERE))
import build_editable_ppt as B  # noqa: E402


def _font(size_px: int):
    """Load a real TrueType face so text metrics are measurable.

    PIL's default bitmap font is ~8px tall no matter what size you ask for,
    which makes a synthetic reference inconsistent with its own declared
    font sizes and turns a correct rebuild into a false failure.
    """
    for p in ("/System/Library/Fonts/Supplemental/Arial Bold.ttf",
              "/System/Library/Fonts/Supplemental/Arial.ttf",
              "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
              "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"):
        if Path(p).exists():
            try:
                return ImageFont.truetype(p, size_px)
            except Exception:
                continue
    return ImageFont.load_default()


def make_reference(path: Path, w: int = 1400, h: int = 900):
    """A deliberately non-16:9, asymmetric layout to exercise full-canvas fit."""
    im = Image.new("RGB", (w, h), "#FDFEFF")
    d = ImageDraw.Draw(im)

    f_title = _font(22)
    f_head = _font(17)
    f_card = _font(15)

    d.rectangle([0, 0, w - 1, 74], fill="#00439E")
    d.text((28, 26), "TITLE BAR SPANS FULL WIDTH", fill="white", font=f_title)

    # left main panel
    d.rounded_rectangle([32, 96, 690, 520], radius=14, outline="#008980",
                        width=3, fill="#FFFFFF")
    d.rectangle([32, 96, 690, 140], fill="#008980")
    d.text((60, 108), "LEFT PANEL HEADER", fill="white", font=f_head)

    # right main panel — deliberately different width
    d.rounded_rectangle([730, 96, 1368, 420], radius=14, outline="#C0392B",
                        width=3, fill="#FFFFFF")
    d.rectangle([730, 96, 1368, 140], fill="#C0392B")
    d.text((760, 108), "RIGHT PANEL HEADER", fill="white", font=f_head)

    # bottom panel — offset, not centered
    d.rounded_rectangle([120, 560, 900, 850], radius=14, outline="#7D3C98",
                        width=3, fill="#FFFFFF")
    d.rectangle([120, 560, 900, 604], fill="#7D3C98")
    d.text((150, 572), "BOTTOM PANEL HEADER", fill="white", font=f_head)

    # inner cards
    for i in range(3):
        x = 60 + i * 210
        d.rounded_rectangle([x, 170, x + 190, 300], radius=8, outline="#888888",
                            width=2, fill="#F4F4F6")
        d.text((x + 20, 226), f"CARD {i + 1}", fill="#333333", font=f_card)

    d.ellipse([760, 170, 860, 270], outline="#00439E", width=4, fill="#E8F0FA")
    d.text((785, 212), "CIRC", fill="#00439E", font=f_card)

    # content reaching the right and bottom edges
    d.rectangle([1330, 150, 1366, 400], fill="#E67E22")
    d.rectangle([130, 830, 890, 848], fill="#16A085")

    im.save(path)
    return w, h


def build_spec(w: int, h: int) -> dict:
    return {
        "page": {"width_px": w, "height_px": h, "background": "#FDFEFF"},
        "frames": [
            {"name": "title_bar", "type": "rectangle", "x": 0, "y": 0,
             "w": w, "h": 74, "fill": "#00439E", "line": "#00439E",
             "line_width_px": 2, "z": 5},
            {"name": "left_main_panel", "type": "rounded_rectangle",
             "x": 32, "y": 96, "w": 658, "h": 424, "fill": "#FFFFFF",
             "line": "#008980", "line_width_px": 3, "corner_radius": 0.05,
             "z": 20},
            {"name": "left_header", "type": "rectangle", "x": 32, "y": 96,
             "w": 658, "h": 44, "fill": "#008980", "line": "#008980",
             "line_width_px": 2, "z": 25},
            {"name": "right_main_panel", "type": "rounded_rectangle",
             "x": 730, "y": 96, "w": 638, "h": 324, "fill": "#FFFFFF",
             "line": "#C0392B", "line_width_px": 3, "corner_radius": 0.05,
             "z": 20},
            {"name": "right_header", "type": "rectangle", "x": 730, "y": 96,
             "w": 638, "h": 44, "fill": "#C0392B", "line": "#C0392B",
             "line_width_px": 2, "z": 25},
            {"name": "bottom_main_panel", "type": "rounded_rectangle",
             "x": 120, "y": 560, "w": 780, "h": 290, "fill": "#FFFFFF",
             "line": "#7D3C98", "line_width_px": 3, "corner_radius": 0.05,
             "z": 20},
            {"name": "bottom_header", "type": "rectangle", "x": 120, "y": 560,
             "w": 780, "h": 44, "fill": "#7D3C98", "line": "#7D3C98",
             "line_width_px": 2, "z": 25},
            {"name": "card_01", "type": "rounded_rectangle", "x": 60, "y": 170,
             "w": 190, "h": 130, "fill": "#F4F4F6", "line": "#888888",
             "line_width_px": 2, "corner_radius": 0.06, "z": 30},
            {"name": "card_02", "type": "rounded_rectangle", "x": 270, "y": 170,
             "w": 190, "h": 130, "fill": "#F4F4F6", "line": "#888888",
             "line_width_px": 2, "corner_radius": 0.06, "z": 30},
            {"name": "card_03", "type": "rounded_rectangle", "x": 480, "y": 170,
             "w": 190, "h": 130, "fill": "#F4F4F6", "line": "#888888",
             "line_width_px": 2, "corner_radius": 0.06, "z": 30},
            {"name": "process_circle_01", "type": "oval", "x": 760, "y": 170,
             "w": 100, "h": 100, "fill": "#E8F0FA", "line": "#00439E",
             "line_width_px": 3, "z": 30},
            {"name": "right_edge_block", "type": "rectangle", "x": 1330,
             "y": 150, "w": 36, "h": 250, "fill": "#E67E22",
             "line": "#E67E22", "line_width_px": 2, "z": 28},
            {"name": "bottom_edge_block", "type": "rectangle", "x": 130,
             "y": 830, "w": 760, "h": 18, "fill": "#16A085",
             "line": "#16A085", "line_width_px": 2, "z": 28},
        ],
        "texts": [
            {"name": "title_text", "text": "TITLE BAR SPANS FULL WIDTH",
             "x": 28, "y": 24, "w": 700, "h": 30, "font": "Microsoft YaHei",
             "font_size_px": 22, "bold": True, "color": "#FFFFFF",
             "align": "left", "vertical_align": "middle",
             "margin_left_px": 0, "margin_right_px": 0,
             "margin_top_px": 0, "margin_bottom_px": 0, "z": 80},
            {"name": "left_header_text", "text": "LEFT PANEL HEADER",
             "x": 60, "y": 108, "w": 400, "h": 22, "font": "Microsoft YaHei",
             "font_size_px": 17, "bold": True, "color": "#FFFFFF",
             "align": "left", "vertical_align": "middle", "z": 82},
            {"name": "right_header_text", "text": "RIGHT PANEL HEADER",
             "x": 760, "y": 108, "w": 400, "h": 22, "font": "Microsoft YaHei",
             "font_size_px": 17, "bold": True, "color": "#FFFFFF",
             "align": "left", "vertical_align": "middle", "z": 82},
            {"name": "bottom_header_text", "text": "BOTTOM PANEL HEADER",
             "x": 150, "y": 572, "w": 400, "h": 22, "font": "Microsoft YaHei",
             "font_size_px": 17, "bold": True, "color": "#FFFFFF",
             "align": "left", "vertical_align": "middle", "z": 82},
            {"name": "card_01_text", "text": "CARD 1", "x": 80, "y": 225,
             "w": 150, "h": 24, "font": "Microsoft YaHei", "font_size_px": 15,
             "color": "#333333", "align": "left", "vertical_align": "middle",
             "z": 82},
            {"name": "card_02_text", "text": "CARD 2", "x": 290, "y": 225,
             "w": 150, "h": 24, "font": "Microsoft YaHei", "font_size_px": 15,
             "color": "#333333", "align": "left", "vertical_align": "middle",
             "z": 82},
            {"name": "card_03_text", "text": "CARD 3", "x": 500, "y": 225,
             "w": 150, "h": 24, "font": "Microsoft YaHei", "font_size_px": 15,
             "color": "#333333", "align": "left", "vertical_align": "middle",
             "z": 82},
            {"name": "process_circle_01_text", "text": "CIRC", "x": 775,
             "y": 208, "w": 70, "h": 24, "font": "Microsoft YaHei",
             "font_size_px": 15, "color": "#00439E", "align": "center",
             "vertical_align": "middle", "z": 82},
        ],
        "icons": [],
        "connectors": [
            {"name": "connector_top_left", "type": "elbow", "x1": 1000,
             "y1": 146, "x2": 300, "y2": 173, "line": "#00439E",
             "line_width_px": 4, "end_arrow": True, "start_arrow": False,
             "layer": "outer", "z": 5},
            {"name": "connector_top_right", "type": "straight", "x1": 1020,
             "y1": 150, "x2": 1290, "y2": 200, "line": "#C0392B",
             "line_width_px": 4, "end_arrow": True, "layer": "outer",
             "z": 5},
            {"name": "connector_center_bidirectional", "type": "straight",
             "x1": 790, "y1": 300, "x2": 1200, "y2": 300, "line": "#7D3C98",
             "line_width_px": 4, "end_arrow": True, "start_arrow": True,
             "layer": "outer", "z": 6},
            {"name": "connector_inner_01", "type": "elbow", "x1": 250,
             "y1": 300, "x2": 700, "y2": 330, "line": "#666666",
             "line_width_px": 3, "end_arrow": True, "layer": "inner",
             "z": 35},
        ],
    }


def main():
    tmp = Path(tempfile.mkdtemp(prefix="pl_ppt_selftest_"))
    results = []

    def check(name, ok, detail=""):
        results.append((name, ok, detail))
        print(f"{'PASS' if ok else 'FAIL'}  {name}" +
              (f"  — {detail}" if detail else ""))

    try:
        ref_path = tmp / "reference.png"
        w, h = make_reference(ref_path)
        check("reference generated", ref_path.exists(), f"{w}x{h}")

        # --- 1. full-canvas mapping: no contain/fit, no letterbox ---
        canvas = B.Canvas(w, h)
        check("slide ratio matches reference",
              abs((canvas.slide_w_in / canvas.slide_h_in) - (w / h)) < 1e-6,
              f"{canvas.slide_w_in:.3f}x{canvas.slide_h_in:.3f}in "
              f"vs {w}x{h}")
        check("X(0) is left edge", canvas.X(0) == 0)
        check("X(REF_W) is right edge",
              abs(canvas.X(w).inches - canvas.slide_w_in) < 1e-9)
        check("Y(REF_H) is bottom edge",
              canvas.Y(h) == canvas.slide_h_emu,
              f"{canvas.Y(h)} == {canvas.slide_h_emu}")

        # --- 2. spec validation ---
        spec = build_spec(w, h)
        spec_path = tmp / "scene_spec.json"
        spec_path.write_text(json.dumps(spec, ensure_ascii=False), "utf-8")
        r = subprocess.run(
            [sys.executable, str(HERE / "validate_spec.py"), str(spec_path),
             "--ref", str(ref_path)], capture_output=True, text=True)
        check("spec validates", r.returncode == 0, r.stdout.strip().splitlines()[-1]
              if r.stdout else r.stderr)

        # --- 3. build ---
        pptx = B.build_presentation(spec, canvas, tmp / "icons",
                                    tmp / "ppt" / "editable_reproduction.pptx")
        check("pptx built", pptx.exists(), f"{pptx.stat().st_size} bytes")

        from pptx import Presentation
        prs = Presentation(str(pptx))
        slide = prs.slides[0]
        check("slide count is 1", len(prs.slides) == 1)
        check("slide size is full-canvas",
              prs.slide_width == canvas.slide_w_emu and
              prs.slide_height == canvas.slide_h_emu,
              f"{prs.slide_width}x{prs.slide_height} EMU")
        # Aspect ratio must match the reference EXACTLY, not approximately.
        check("slide aspect ratio error is 0",
              prs.slide_width * h == prs.slide_height * w,
              f"{prs.slide_width * h} vs {prs.slide_height * w}")

        # --- 4. no reference image embedded ---
        check("no picture shapes in slide",
              all(sh.shape_type != 13 for sh in slide.shapes),
              f"{len(slide.shapes)} shapes, all native")

        # --- 4b. no theme shadow on any shape ---
        import zipfile
        with zipfile.ZipFile(pptx) as zf:
            xml = zf.read("ppt/slides/slide1.xml").decode()
        check("no theme <p:style> shadow references",
              xml.count("<p:style>") == 0,
              f"{xml.count('<p:style>')} style refs")
        check("no outerShdw effects", xml.count("outerShdw") == 0)

        # --- 5. native connectors, no freeform ---
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        conns = [sh for sh in slide.shapes
                 if sh.shape_type == MSO_SHAPE_TYPE.LINE]
        freeform = [sh for sh in slide.shapes
                    if sh.shape_type == MSO_SHAPE_TYPE.FREEFORM]
        check("connectors are native lines", len(conns) == 4,
              f"{len(conns)} connectors")
        check("zero freeform shapes", len(freeform) == 0)

        # --- 6. out of bounds ---
        oob = B.find_out_of_bounds(slide, prs)
        check("no out-of-bounds shapes", len(oob) == 0, str(oob))

        # --- 7. media check ---
        media = B.check_media(pptx)
        check("no huge media", len(media["huge_media"]) == 0,
              f"{media['media_count']} media files")

        # --- 8. render + compare ---
        png = B.render_pptx(pptx, tmp / "preview")
        if not png:
            check("ppt rendered", False,
                  "LibreOffice unavailable - install with "
                  "`brew install --cask libreoffice`")
        else:
            check("ppt rendered", True, str(png))
            reference, _, _ = B.load_reference(ref_path)
            render = Image.open(png).convert("RGB")
            files = B.make_comparisons(reference, render, tmp / "preview")
            for k, v in files.items():
                check(f"{k} generated", Path(v).exists(), v)
            m = B.measure_page(reference, render)
            print("     " + json.dumps(m, ensure_ascii=False))
            check("occupancy error X <= 0.005",
                  m["occupancy_error_x"] <= 0.005, str(m["occupancy_error_x"]))
            check("occupancy error Y <= 0.005",
                  m["occupancy_error_y"] <= 0.005, str(m["occupancy_error_y"]))
            check("center error X <= 0.003",
                  m["center_error_x"] <= 0.003, str(m["center_error_x"]))
            check("center error Y <= 0.003",
                  m["center_error_y"] <= 0.003, str(m["center_error_y"]))

        # --- 8a. panel borders must land on the same pixels, not near them ---
        if png:
            reference3, _, _ = B.load_reference(ref_path)
            render3 = Image.open(png).convert("RGB").resize(
                reference3.size, Image.Resampling.LANCZOS)
            import numpy as np

            def border_col(im, color, x_lo, x_hi, y0, y1):
                a = np.array(im.convert("RGB"))[y0:y1, x_lo:x_hi]
                d = np.abs(a - np.array(color)).sum(axis=2)
                ys, xs = np.where(d < 40)
                return int(xs.min() + x_lo) if len(xs) else None

            c_teal = (0, 137, 128)
            rx = border_col(reference3, c_teal, 10, 60, 200, 400)
            kx = border_col(render3, c_teal, 10, 60, 200, 400)
            check("teal panel left border on the same pixel",
                  rx is not None and kx is not None and abs(rx - kx) <= 1,
                  f"ref x={rx} ren x={kx}")

        # --- 8b. text height must match the reference, not just fit ---
        if png:
            reference2, _, _ = B.load_reference(ref_path)
            render2 = Image.open(png).convert("RGB").resize(
                reference2.size, Image.Resampling.LANCZOS)
            import numpy as np

            def title_text_height(im):
                g = np.array(im.convert("L"))[0:74, 28:700]
                ys, xs = np.where(g > 200)
                return (ys.max() - ys.min() + 1) if len(ys) else 0

            h_ref = title_text_height(reference2)
            h_ren = title_text_height(render2)
            check("title text height matches reference",
                  abs(h_ren - h_ref) <= max(3, h_ref * 0.35),
                  f"ref={h_ref}px ren={h_ren}px")

        # --- 9. icon pipeline ---
        svg = tmp / "icons" / "svg" / "trophy_blue.svg"
        rr = subprocess.run(
            [sys.executable, str(HERE / "redraw_icon.py"), "--kind", "trophy",
             "--out", str(svg), "--color", "#1B7FBF"],
            capture_output=True, text=True)
        check("icon svg generated", rr.returncode == 0 and svg.exists())
        png128 = tmp / "icons" / "png_128" / "trophy_blue.png"
        ok = B.svg_to_png_128(svg, png128)
        check("svg rendered to png", ok)
        if ok:
            with Image.open(png128) as im:
                check("png is 128x128 RGBA",
                      im.size == (128, 128) and im.mode == "RGBA",
                      f"{im.size} {im.mode}")
                import numpy as np
                a = np.array(im)[:, :, 3]
                corners = [a[0, 0], a[0, 127], a[127, 0], a[127, 127]]
                check("png corners are transparent",
                      all(c == 0 for c in corners), str(corners))
            fx, fy, fw, fh = B.alpha_bbox_fraction(png128)
            check("alpha bbox is a real region", 0 < fw <= 1 and 0 < fh <= 1,
                  f"fx={fx:.3f} fy={fy:.3f} fw={fw:.3f} fh={fh:.3f}")

        # --- 10. manifest ---
        mm = subprocess.run(
            [sys.executable, str(HERE / "make_manifest.py"),
             "--project", str(tmp), "--reference", str(ref_path)],
            capture_output=True, text=True)
        manifest = tmp / "manifest.txt"
        check("manifest written", mm.returncode == 0 and manifest.exists())
        if manifest.exists():
            txt = manifest.read_text(encoding="utf-8")
            for key in ("Full-canvas mapping used", "Freeform arrows used",
                        "Reference image used as background",
                        "PPTX file size"):
                check(f"manifest has '{key}'", key in txt)

        # --- 11. packaging ---
        pp = subprocess.run(
            [sys.executable, str(HERE / "package_outputs.py"),
             "--project", str(tmp), "--out", str(tmp / "dist")],
            capture_output=True, text=True)
        check("packaging succeeded", pp.returncode == 0, pp.stderr.strip()[:120])
        check("project zip created",
              (tmp / "dist" / f"{tmp.name}.zip").exists())
        check("icons zip created", (tmp / "dist" / "icons_assets.zip").exists())

    finally:
        passed = sum(1 for _, ok, _ in results if ok)
        total = len(results)
        print(f"\n{passed}/{total} PASS")
        if passed != total:
            print("SELFTEST FAILED")
            sys.exit(1)
        print("SELFTEST OK")
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == "__main__":
    main()
